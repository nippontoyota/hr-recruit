from uuid import UUID
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.deps import require_roles
from app.core.access import get_candidate_for_user
from app.models.candidate import Candidate
from app.models.activity_log import ActivityLog
from app.models.enums import PipelineStage, UserRole, ActivityType
from app.models.stage_history import StageHistory
from app.models.user import User
from app.models.communication import Communication
from app.models.enums import CommunicationType, CommunicationDirection, CommunicationStatus
from app.schemas.candidate import CandidateOut, DocumentOut, StageChange, StageHistoryOut, ActivityLogOut, VisitScheduleUpdate, WhatsAppInviteCreate, CandidateDepartmentUpdate
from app.services.workflow import transition
from app.services.doubletick import send_template, DoubleTickError, friendly_doubletick_error

from app.services.candidate_service import (
    to_candidate_out,
)
from app.api.v1.candidates_core import (
    _issue_pre_form as issue_pre_form,
    _store_whatsapp_invite as store_whatsapp_invite,
)
from app.services.document_service import (
    get_resume_document,
    document_out,
    save_resume_for_candidate,
    save_photo_for_candidate,
    resume_candidate_ids,
)

router = APIRouter(prefix="/candidates", tags=["candidates"])

@router.post("/{id}/resume", response_model=DocumentOut, status_code=201)
async def upload_resume(
    id: UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    row = get_candidate_for_user(db, id, user)
    return await save_resume_for_candidate(db, row, file, uploaded_by_user_id=user.id)


@router.post("/{id}/photo", response_model=dict, status_code=201)
async def upload_photo(
    id: UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    row = db.get(Candidate, id)
    if not row:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    return await save_photo_for_candidate(db, row, file)


@router.get("/{id}/resume", response_model=DocumentOut)
def get_resume(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    get_candidate_for_user(db, id, user)
    doc = get_resume_document(db, id)
    if not doc:
        raise HTTPException(status_code=404, detail="Resume not found.")
    return document_out(doc)


@router.post("/{id}/transition", response_model=CandidateOut)
def transition_stage(
    id: UUID,
    body: StageChange,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    row = get_candidate_for_user(db, id, user)
    updated = transition(
        db=db,
        candidate=row,
        target_stage=body.to_stage,
        user=user,
        remarks=body.remarks
    )
    db.commit()
    db.refresh(updated)
    return to_candidate_out(updated, id in resume_candidate_ids(db, [id]))


class UnholdRequest(BaseModel):
    remarks: str | None = None


@router.post("/{id}/unhold", response_model=CandidateOut)
def unhold_candidate(
    id: UUID,
    body: UnholdRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    candidate = get_candidate_for_user(db, id, user)
    last_hold = db.scalars(
        select(StageHistory)
        .where(StageHistory.candidate_id == id, StageHistory.to_stage == PipelineStage.ON_HOLD)
        .order_by(StageHistory.created_at.desc())
    ).first()

    if not last_hold:
        raise HTTPException(status_code=400, detail="No previous ON_HOLD transition found for candidate.")

    target_stage = last_hold.from_stage
    if not target_stage:
        raise HTTPException(status_code=400, detail="Cannot determine previous stage to resume to.")

    remarks = body.remarks or "Resumed from On Hold"

    updated = transition(db=db, candidate=candidate, target_stage=target_stage, user=user, remarks=remarks)
    db.commit()
    db.refresh(updated)
    return to_candidate_out(updated, id in resume_candidate_ids(db, [id]))


@router.get("/{id}/stage-history", response_model=list[StageHistoryOut])
def history(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    get_candidate_for_user(db, id, user)
    return list(db.scalars(select(StageHistory).where(StageHistory.candidate_id == id).order_by(StageHistory.created_at)))

@router.get("/{id}/activity-logs", response_model=list[ActivityLogOut])
def get_activity_logs(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    get_candidate_for_user(db, id, user)
    return list(db.scalars(select(ActivityLog).where(ActivityLog.candidate_id == id).order_by(ActivityLog.created_at.desc())))

@router.patch("/{id}/visit-schedule", response_model=CandidateOut)
def update_visit_schedule(
    id: UUID,
    body: VisitScheduleUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    candidate = get_candidate_for_user(db, id, user)
    
    if body.visit_branch is not None:
        candidate.visit_branch = body.visit_branch
    if body.visit_date is not None:
        candidate.visit_date = body.visit_date
    if body.visit_time is not None:
        candidate.visit_time = body.visit_time
    if body.visit_maps_link is not None:
        candidate.visit_maps_link = body.visit_maps_link
    if body.visit_instructions is not None:
        candidate.visit_instructions = body.visit_instructions

    log = ActivityLog(
        candidate_id=id,
        activity_type=ActivityType.SYSTEM,
        title="Visit Schedule Updated",
        description=f"Visit schedule updated. Branch: {candidate.visit_branch}, Date: {candidate.visit_date}",
        created_by_user_id=user.id,
    )
    db.add(log)
    db.commit()
    db.refresh(candidate)
    return to_candidate_out(candidate, id in resume_candidate_ids(db, [id]))


ALLOWED_DEPARTMENTS = {
    "Sales",
    "Service",
    "Customer Service",
    "Insurance",
    "Call Center",
    "HR",
}


@router.patch("/{id}/department", response_model=CandidateOut)
def update_candidate_department(
    id: UUID,
    body: CandidateDepartmentUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    """Change the department and/or role the candidate is being considered for."""
    candidate = get_candidate_for_user(db, id, user)
    department = body.department.strip()
    if department not in ALLOWED_DEPARTMENTS:
        raise HTTPException(
            status_code=400,
            detail=f"Department must be one of: {', '.join(sorted(ALLOWED_DEPARTMENTS))}",
        )

    role = (body.position_applied_for or "").strip() or department
    previous_dept = candidate.department
    previous_role = candidate.position_applied_for

    if previous_dept == department and previous_role == role:
        return to_candidate_out(candidate, id in resume_candidate_ids(db, [id]))

    candidate.department = department
    candidate.position_applied_for = role
    db.add(
        ActivityLog(
            candidate_id=id,
            activity_type=ActivityType.SYSTEM,
            title="Consideration Updated",
            description=(
                f"Considering for changed from {previous_dept or '—'} / {previous_role or '—'} "
                f"to {department} / {role}."
            ),
            created_by_user_id=user.id,
        )
    )
    db.commit()
    db.refresh(candidate)
    return to_candidate_out(candidate, id in resume_candidate_ids(db, [id]))

@router.post("/{id}/pre-form/send", response_model=CandidateOut)
def send_pre_form(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    row = get_candidate_for_user(db, id, user)
    issue_pre_form(db, row, user)
    db.commit()
    db.refresh(row)
    return to_candidate_out(row, id in resume_candidate_ids(db, [id]))

@router.post("/{id}/whatsapp-invite")
def send_whatsapp_invite(
    id: UUID,
    body: WhatsAppInviteCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR, UserRole.LOCAL_HR)),
):
    candidate = get_candidate_for_user(db, id, user)
    vars_map = body.variables or {}
    if not (vars_map.get("visitDate") or "").strip():
        raise HTTPException(status_code=400, detail="Visit date is required before sending.")
    if not (vars_map.get("branchName") or "").strip() or not (vars_map.get("mapsLink") or "").strip():
        raise HTTPException(status_code=400, detail="Location and maps link are required before sending.")

    # Order must match DoubleTick template nippon_interview_call_letter {{1}}…{{9}}
    DOUBLETICK_VARIABLE_KEYS = [
        "candidateName",
        "position",
        "visitDate",
        "branchName",
        "formLink",
        "arrivalTime",
        "extraInstructions",
        "mapsLink",
        "recruiterName",
    ]

    placeholders = []
    for key in DOUBLETICK_VARIABLE_KEYS:
        placeholders.append(vars_map.get(key, ""))

    try:
        res = send_template(
            to_phone=candidate.phone,
            template_name="nippon_interview_call_letter",
            placeholders=placeholders,
        )
        external_message_id = None
        messages = res.get("messages", [])
        if messages:
            external_message_id = messages[0].get("id")

        status = CommunicationStatus.SENT
        err_msg = None
    except DoubleTickError as e:
        status = CommunicationStatus.FAILED
        err_msg = e.user_message
    except Exception as e:
        status = CommunicationStatus.FAILED
        err_msg = friendly_doubletick_error(str(e))

    extra = (vars_map.get("extraInstructions") or "").strip() or (
        "Meeting Point – Floor 3rd – Sales Training Room / HR Department\n"
        "Touch Point 1 – Sreehari (HRD) 8606986060\n"
        "Touch Point 2 – Mathew (HRD) 9544286099"
    )
    content_lines = [
        f"Dear {vars_map.get('candidateName', '')},",
        "",
        '"Greetings from Nippon HRD"',
        "",
        (
            f"This is to inform you that, pertaining to your application for "
            f"*{vars_map.get('position', '')}*, we have scheduled a direct interview on "
            f"*{vars_map.get('visitDate', '')}* at Nippon Toyota, *{vars_map.get('branchName', '')}*. "
            "Please bring an updated bio-data and a passport size photo."
        ),
        "",
        "Also complete the Job Application Form using the link below without fail:",
        vars_map.get("formLink", ""),
        "",
        f"Reporting Time – *{vars_map.get('arrivalTime', '')}*",
        "Dress Code – Formal Wear with Proper Grooming (Mandatory)",
        extra,
        "",
        "Location Link –",
        vars_map.get("mapsLink", ""),
        "",
        "Regards",
        vars_map.get("recruiterName", ""),
        "Talent Acquisition Team",
        "Nippon Toyota",
    ]
    full_content = "\n".join(content_lines)

    comm = Communication(
        candidate_id=id,
        type=CommunicationType.WHATSAPP,
        direction=CommunicationDirection.OUTGOING,
        status=status,
        content_preview=full_content[:255],
        external_message_id=external_message_id,
        created_by=user.id
    )
    db.add(comm)

    activity_desc = f"Template: nippon_interview_call_letter. Status: {status.value}."
    if err_msg:
        activity_desc += f" Error: {err_msg}"
        
    log = ActivityLog(
        candidate_id=id,
        activity_type=ActivityType.WHATSAPP,
        title="WhatsApp Invite Sent" if status == CommunicationStatus.SENT else "WhatsApp Invite Failed",
        description=activity_desc,
        created_by_user_id=user.id,
    )
    db.add(log)
    db.commit()
    
    if status == CommunicationStatus.FAILED:
        raise HTTPException(status_code=400, detail=err_msg)
        
    store_whatsapp_invite(db, candidate, user)
    db.commit()
    return {"status": "success", "message_id": external_message_id}


from app.api.v1.pdf import generate_offer_letter_pdf
from app.services.email import send_email_with_pdf

class SendOfferLetterRequest(BaseModel):
    pass

@router.post("/{id}/offer-letter/send", response_model=CandidateOut)
def send_offer_letter(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR)),
):
    row = get_candidate_for_user(db, id, user)
    
    if row.current_stage not in (PipelineStage.FINAL_APPROVAL, PipelineStage.HIRED):
        raise HTTPException(status_code=400, detail="Offer letter can only be sent in Final Approval or Hired stage.")
        
    if not row.email:
        raise HTTPException(status_code=400, detail="Candidate does not have an email address.")
        
    payload = {
        "candidate": {
            "full_name": row.full_name,
            "position_applied_for": row.department,
            "salary_data": row.salary_data
        }
    }
    pdf_bytes = generate_offer_letter_pdf(payload)
    
    subject = f"Offer of Employment - {row.department} at Nippon Toyota"
    body_html = f"""
    <html>
        <body>
            <p>Dear {row.full_name},</p>
            <p>We are delighted to offer you the position of <strong>{row.department}</strong> at Nippon Toyota.</p>
            <p>Please find your official offer letter attached as a PDF.</p>
            <p>We look forward to welcoming you to the team!</p>
            <br/>
            <p>Best regards,<br/>Human Resources<br/>Nippon Toyota</p>
        </body>
    </html>
    """
    
    send_email_with_pdf(
        to_email=row.email,
        subject=subject,
        body_html=body_html,
        pdf_bytes=bytes(pdf_bytes),
        pdf_filename="OfferLetter_NipponToyota.pdf"
    )
    
    comm = Communication(
        candidate_id=id,
        type=CommunicationType.EMAIL,
        direction=CommunicationDirection.OUTGOING,
        status=CommunicationStatus.SENT,
        subject=subject,
        content_preview="Sent Offer Letter Email with PDF Attachment.",
        created_by=user.id
    )
    db.add(comm)
    
    log = ActivityLog(
        candidate_id=id,
        activity_type=ActivityType.EMAIL,
        title="Sent Offer Letter",
        description="Generated and sent offer letter PDF via email.",
        created_by_user_id=user.id
    )
    db.add(log)
    
    row.offer_status = "SENT"
    
    db.commit()
    db.refresh(row)
    
    return to_candidate_out(row, id in resume_candidate_ids(db, [id]))

@router.post("/bulk-salary", status_code=200)
async def upload_bulk_salary(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.ADMIN, UserRole.HO_HR)),
):
    import openpyxl
    import io
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Only Excel files (.xlsx, .xls) are supported.")
    
    try:
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        
        headers = []
        updated_count = 0
        not_found_count = 0
        
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(cell).strip() if cell else f"col_{j}" for j, cell in enumerate(row)]
                continue
                
            if not any(row):
                continue
                
            row_data = dict(zip(headers, row))
            candidate_id = row_data.get("Candidate ID") or row_data.get("candidate_id")
            email = row_data.get("Email") or row_data.get("email")
            
            query = select(Candidate)
            if candidate_id:
                query = query.where(Candidate.candidate_id == str(candidate_id))
            elif email:
                query = query.where(Candidate.email == str(email))
            else:
                continue
                
            candidate = db.scalar(query)
            if candidate:
                candidate.salary_data = row_data
                updated_count += 1
            else:
                not_found_count += 1
                
        db.commit()
        return {
            "message": f"Successfully updated {updated_count} candidates.",
            "updated_count": updated_count,
            "not_found_count": not_found_count
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to process excel file: {str(e)}")

@router.post("/{id}/send-to-ho", response_model=CandidateOut)
def send_to_ho(
    id: UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(UserRole.LOCAL_HR, UserRole.ADMIN)),
):
    row = get_candidate_for_user(db, id, user)
    
    if row.current_stage != PipelineStage.APPLICATION:
        raise HTTPException(status_code=400, detail="Candidate must be in APPLICATION stage to be sent to Head Office.")
        
    updated = transition(
        db=db,
        candidate=row,
        target_stage=PipelineStage.SENT_TO_HO,
        user=user,
        remarks="Application transferred to Head Office"
    )
    db.commit()
    db.refresh(updated)
    return to_candidate_out(updated, id in resume_candidate_ids(db, [id]))
